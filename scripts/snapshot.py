"""Registro modifiche AFFITTI.

Scarica AFFITTI.xlsx da Dropbox, legge l'area di stampa del foglio AFFITTI,
la confronta con l'ultima versione salvata e, se e' cambiata, salva:

  data/versioni/<id>.json   copia completa della tabella (la "foto")
  data/indice.json          elenco versioni + eventi (aggiunte/rimosse/modificate)
  data/stato.json           ultimo controllo (aggiornato al massimo una volta al giorno
                            se non ci sono modifiche, per non creare commit inutili)

Uso:
  python scripts/snapshot.py                 # scarica da DROPBOX_URL (o link di default)
  python scripts/snapshot.py --file a.xlsx   # usa un file locale (prove)
"""

import argparse
import hashlib
import io
import json
import os
import re
import sys
import time
import urllib.request
from datetime import datetime, timezone

import openpyxl
from openpyxl.utils import column_index_from_string, range_boundaries

DEFAULT_URL = ("https://www.dropbox.com/scl/fi/q3ewoodokhsrd8g11l915/AFFITTI.xlsx"
               "?rlkey=47altne9kdaw9y5h8bm57lun5&st=qjbkqviq&dl=0")
SHEET = "AFFITTI"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.environ.get("AFFITTI_DATA") or os.path.join(ROOT, "data")
CAMPI = ["pt", "sopp", "cort", "tot", "canone"]


# ---------------------------------------------------------------- download
def direct_link(url):
    u = url.strip().replace("www.dropbox.com", "dl.dropboxusercontent.com")
    if "dl=0" in u:
        u = u.replace("dl=0", "dl=1")
    elif "dl=1" not in u:
        u += ("&" if "?" in u else "?") + "dl=1"
    return u


def download(url, tentativi=3):
    err = None
    for i in range(tentativi):
        try:
            req = urllib.request.Request(direct_link(url), headers={"User-Agent": "affitti-registro"})
            with urllib.request.urlopen(req, timeout=60) as r:
                data = r.read()
            if data[:2] != b"PK":
                raise ValueError("risposta non e' un file xlsx")
            return data
        except Exception as e:  # noqa: BLE001
            err = e
            time.sleep(10 * (i + 1))
    raise SystemExit("Download fallito: %s" % err)


# ---------------------------------------------------------------- lettura
def valore(v):
    if isinstance(v, datetime):
        return {"data": v.strftime("%Y-%m-%dT%H:%M")}
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip()
    return v


def area_stampa(ws):
    pa = ws.print_area
    if isinstance(pa, (list, tuple)):
        pa = pa[0] if pa else None
    if pa:
        ref = pa.split("!")[-1].replace("$", "")
        c0, r0, c1, r1 = range_boundaries(ref)
        return r0, r1, c0, c1
    # ripiego: colonne A-F fino alla prima riga vuota dopo l'intestazione
    r1 = 2
    while r1 < ws.max_row and any(ws.cell(r1 + 1, c).value is not None for c in range(1, 7)):
        r1 += 1
    return 1, r1, 1, 6


def colonne_nascoste(ws):
    nascoste = set()
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            lo = dim.min or column_index_from_string(key)
            hi = dim.max or lo
            nascoste.update(range(lo, hi + 1))
    return nascoste


def leggi_tabella(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit("Foglio %r non trovato" % SHEET)
    ws = wb[SHEET]
    r0, r1, c0, c1 = area_stampa(ws)
    hid_c = colonne_nascoste(ws)
    righe_vis = [r for r in range(r0, r1 + 1) if not ws.row_dimensions[r].hidden]
    col_vis = [c for c in range(c0, c1 + 1) if c not in hid_c]

    griglia = [[valore(ws.cell(r, c).value) for c in col_vis] for r in righe_vis]

    unioni = []
    for m in ws.merged_cells.ranges:
        rr = [i for i, r in enumerate(righe_vis) if m.min_row <= r <= m.max_row]
        cc = [i for i, c in enumerate(col_vis) if m.min_col <= c <= m.max_col]
        if rr and cc:
            unioni.append([rr[0], cc[0], rr[-1], cc[-1]])

    # data/ora di Excel (cella con =NOW() nell'intestazione)
    excel_ora = None
    for v in griglia[0] if griglia else []:
        if isinstance(v, dict) and "data" in v:
            excel_ora = v["data"]

    return {"griglia": griglia, "unioni": unioni}, excel_ora


# ---------------------------------------------------------------- unita
# errori di battitura visti nelle schede: stessa unita'
ALIAS = {"stocaggio": "stoccaggio", "montacarich": "montacarichi", "motacarichi": "montacarichi",
         "essecatura": "essiccatura"}


def chiave(nome):
    """Chiave che riconosce la stessa unita': minuscolo, senza emoji/simboli,
    senza spazi (super condor 9 = supercondor 9), refusi noti corretti."""
    k = nome.lower()
    k = re.sub(r"[^\w\s/.\-]", " ", k)          # via emoji e simboli
    k = " ".join(ALIAS.get(p, p) for p in k.split())
    return re.sub(r"\s+", "", k)


def canon_sezione(testo):
    """Nome uniforme della via/sezione (senza numeri di telefono)."""
    t = str(testo or "").split("\n")[0].strip()
    u = t.upper()
    if "TUCIDIDE" in u:
        return "VIA TUCIDIDE"
    if "MALAGA" in u:
        return "VIA MALAGA"
    if "CASCATA" in u or "ROZZANO" in u:
        return "CASCATA (Rozzano)"
    return re.sub(r"\s*[\d/ ]{6,}$", "", t).strip() or "?"


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def estrai_unita(griglia):
    """Riga 1 = intestazione, riga 2 = etichette (indirizzo nella prima cella).
    Poi: riga senza numeri con la prima parola MAIUSCOLA = nuova sezione
    (es. VIA MALAGA, CASCATA via ...);
    altrimenti e' un'unita' (anche se senza mq/canone)."""
    if len(griglia) < 2:
        return []
    sezione = canon_sezione(griglia[1][0])
    unita, viste = [], {}
    for i, riga in enumerate(griglia[2:], start=2):
        nome = riga[0]
        if not isinstance(nome, str) or not nome.strip():
            continue
        valori = [num(v) for v in riga[1:]]
        primo = re.sub(r"[^A-Za-zÀ-ÿ]", "", nome.split()[0])
        if all(v is None for v in valori) and len(primo) >= 2 and primo == primo.upper():
            sezione = canon_sezione(nome)
            continue
        dati = dict(zip(CAMPI, (valori + [None] * len(CAMPI))[:len(CAMPI)]))
        if dati["tot"] is None and any(dati[c] is not None for c in ("pt", "sopp", "cort")):
            dati["tot"] = sum(dati[c] or 0 for c in ("pt", "sopp", "cort"))
        k = chiave(nome)
        viste[k] = viste.get(k, 0) + 1
        if viste[k] > 1:
            k = "%s #%d" % (k, viste[k])
        unita.append({"key": k, "nome": re.sub(r"\s+", " ", nome).strip(), "sezione": sezione,
                      **dati, "riga": i})
    return unita


def confronta(prima, dopo):
    p = {u["key"]: u for u in prima}
    d = {u["key"]: u for u in dopo}
    eventi = []
    for u in dopo:
        if u["key"] not in p:
            eventi.append({"tipo": "aggiunta", **u})
            continue
        old = p[u["key"]]
        campi = {c: [old.get(c), u.get(c)] for c in CAMPI + ["nome", "sezione"] if old.get(c) != u.get(c)}
        if campi:
            eventi.append({"tipo": "modificata", **u, "campi": campi})
    for u in prima:
        if u["key"] not in d:
            eventi.append({"tipo": "rimossa", **u})
    return eventi


# ---------------------------------------------------------------- salvataggio
def leggi_json(path, default):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return default


def scrivi_json(path, obj, compatto=False):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        if compatto:
            json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
        else:
            json.dump(obj, f, ensure_ascii=False, indent=1)
        f.write("\n")


def impronta(tabella, unita):
    # la cella =NOW() cambia a ogni salvataggio: esclusa dal confronto
    g = [[None if isinstance(v, dict) else v for v in r] for r in tabella["griglia"]]
    s = json.dumps({"g": g, "m": tabella["unioni"], "u": unita}, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(s.encode("utf-8")).hexdigest()[:16]


def totali(unita):
    return {
        "n_unita": len(unita),
        "mq_tot": sum(u["tot"] or 0 for u in unita),
        "canone_tot": sum(u["canone"] or 0 for u in unita),
    }


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--file")
    ap.add_argument("--ora", help="data/ora UTC simulata, per le prove (YYYY-MM-DDTHH:MM:SS)")
    args = ap.parse_args()

    if args.file:
        with open(args.file, "rb") as f:
            xlsx = f.read()
    else:
        xlsx = download(os.environ.get("DROPBOX_URL") or DEFAULT_URL)

    adesso = (datetime.fromisoformat(args.ora).replace(tzinfo=timezone.utc) if args.ora
              else datetime.now(timezone.utc)).replace(microsecond=0)
    tabella, excel_ora = leggi_tabella(xlsx)
    unita = estrai_unita(tabella["griglia"])
    imp = impronta(tabella, unita)

    indice = leggi_json(os.path.join(DATA, "indice.json"), {"versioni": []})
    stato_path = os.path.join(DATA, "stato.json")
    stato = leggi_json(stato_path, {})
    rilevato = adesso.isoformat().replace("+00:00", "Z")

    ultima = indice["versioni"][-1] if indice["versioni"] else None
    if ultima and ultima["impronta"] == imp:
        # nessuna modifica: aggiorna stato.json solo una volta al giorno
        if (stato.get("ultimo_controllo") or "")[:10] != rilevato[:10]:
            stato["ultimo_controllo"] = rilevato
            scrivi_json(stato_path, stato)
            print("Nessuna modifica (stato giornaliero aggiornato).")
        else:
            print("Nessuna modifica.")
        return

    if ultima:
        prec = leggi_json(os.path.join(DATA, "versioni", ultima["id"] + ".json"), {})
        eventi = confronta(prec.get("unita", []), unita)
    else:
        eventi = [{"tipo": "iniziale", **u} for u in unita]

    vid = adesso.strftime("%Y%m%d-%H%M%S")
    meta = {"id": vid, "rilevato": rilevato, "excel": excel_ora, "impronta": imp,
            **totali(unita), "eventi": eventi}
    scrivi_json(os.path.join(DATA, "versioni", vid + ".json"),
                {"id": vid, "rilevato": rilevato, "excel": excel_ora, **tabella, "unita": unita},
                compatto=True)
    indice["versioni"].append(meta)
    scrivi_json(os.path.join(DATA, "indice.json"), indice)
    stato.update({"ultimo_controllo": rilevato, "ultima_versione": vid})
    scrivi_json(stato_path, stato)
    print("Nuova versione %s: %d eventi" % (vid, len(eventi)))


if __name__ == "__main__":
    main()
